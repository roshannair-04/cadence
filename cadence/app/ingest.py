"""Read PDFs / PowerPoints and split them into study sections.

A "unit" is one PDF page or one slide. Units are cleaned, given a title, and
merged until each section has enough text to study and quiz on.
"""
import re
from pathlib import Path

from .config import MIN_SECTION_CHARS, SKIP_PATTERNS

BOILERPLATE_TITLE = re.compile(
    r"^(learning outcomes?|session plan|agenda|references?|thank you|questions\??|"
    r"course (details|plan|outcomes?)|text ?books?|contents?|outline|summary of sessions)\b", re.I)
BOILERPLATE_TEXT = re.compile(r"course coordinator|department:|university:", re.I)


def _clean(s: str) -> str:
    s = s.replace(" ", " ").replace("•", "•").replace("", "•").replace("​", "")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"[ \t]*\n[ \t]*", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


LABEL_LINE = re.compile(r"^[A-Z0-9][^:|]{1,45}:\s")
BULLET_START = re.compile(r"^\s*(?:[•\-–—*▪●○◦]|\d{1,2}[.)])\s*")


def unwrap(text: str) -> str:
    """PDF slide text is hard-wrapped at ~60-90 characters; glue wrapped lines back together."""
    out = []
    for line in text.split("\n"):
        cur = line.strip()
        if not cur:
            out.append("")
            continue
        if len(cur) <= 2 and not cur.isdigit() and not BULLET_START.match(cur):
            continue                      # stray icon glyphs from slide graphics
        if out and out[-1]:
            prev = out[-1]
            starts_new = bool(BULLET_START.match(cur)) or bool(LABEL_LINE.match(cur)) or " | " in cur
            prev_done = prev.endswith((".", "!", "?", ":", ";")) and not cur[:1].islower()
            if not starts_new and not prev_done and len(prev) >= 28 and " | " not in prev:
                out[-1] = prev + " " + cur
                continue
            if cur[:1].islower() and not starts_new and len(prev) >= 12 and " | " not in prev:
                out[-1] = prev + " " + cur
                continue
        out.append(cur)
    return "\n".join(out)


def read_pdf(path: Path):
    from pypdf import PdfReader
    units = []
    for i, page in enumerate(PdfReader(str(path)).pages):
        txt = _clean(unwrap(page.extract_text() or ""))
        if txt:
            units.append({"ref": f"p.{i + 1}", "text": txt, "kind": "pdf"})
    return units


def read_pptx(path: Path):
    from pptx import Presentation
    units = []
    for i, slide in enumerate(Presentation(str(path)).slides):
        parts = []
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r_i, row in enumerate(sh.table.rows):
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    parts.append(" | ".join(cells))
        # drop bare slide numbers
        parts = [p for p in parts if not re.fullmatch(r"\d{1,3}", p)]
        parts = [p for p in parts if not re.match(r"^(AIM4141\b.*|Sessions?\s+\d+(\s*[–-]\s*\d+)?\s*(•.*)?)$", p.strip())]
        txt = _clean("\n".join(parts))
        if txt:
            units.append({"ref": f"slide {i + 1}", "text": txt, "kind": "pptx"})
    return units


def read_text(path: Path):
    txt = _clean(path.read_text(errors="ignore"))
    chunks = [c for c in re.split(r"\n\s*\n", txt) if c.strip()]
    return [{"ref": f"part {i + 1}", "text": c, "kind": "txt"} for i, c in enumerate(chunks)]


def guess_title(unit) -> str:
    lines = [l.strip(" •-\t") for l in unit["text"].split("\n") if len(l.strip(" •-\t")) > 2]
    if not lines:
        return "Untitled"
    if unit["kind"] == "pdf" and len(lines) >= 3:
        last = lines[-1]
        # Footer-style titles: short, no sentence punctuation, not a bullet-ish fact
        if len(last) <= 70 and not last.endswith((".", ",", ";")) and len(last.split()) >= 2:
            if not lines[0].isupper() or len(lines[0]) > 40:
                return last
    # Titles broken over several short lines ("What is Human / Computer / Interaction?")
    title = lines[0]
    if len(title.split()) == 1 and len(lines) > 1 and len(lines[1]) <= 14 and not lines[1].endswith((".", ":")):
        title += " " + lines[1]                       # one-word heading wrapped onto two lines
    elif re.match(r"(what|why|how|who|when|where)\b", title, re.I) and not title.endswith("?"):
        for nxt in lines[1:4]:
            if len(nxt) > 26:
                break
            title += " " + nxt
            if title.endswith("?"):
                break
    title = re.sub(r"\s*Sessions?\s+\d+(\s*[–-]\s*\d+)?.*$", "", title).strip()
    return title[:90]


def tidy_title(t: str) -> str:
    """Shorten over-long headings (e.g. 'Case Question 6: Conduct a ...') to their label."""
    t = t.strip()
    if len(t) > 78 or t.endswith("."):
        head = t.split(":")[0].strip()
        if 2 <= len(head.split()) <= 6 and len(head) >= 6:
            return head
    return t


def weak_title(t: str) -> bool:
    t = t.strip()
    return len(t) < 5 or len(t) > 78 or t.endswith((".", ",")) or len(t.split()) > 11


def sentence_case(t: str) -> str:
    return t[:1].upper() + t[1:] if t else t


def strip_title(unit, title: str) -> str:
    t = unit["text"]
    flat = re.sub(r"\s+", " ", title)
    out = []
    removed = False
    buf = ""
    for line in t.split("\n"):
        if not removed:
            cand = re.sub(r"\s+", " ", (buf + " " + line).strip(" •-"))
            if flat.startswith(cand) or cand == flat:
                buf = cand
                if cand == flat:
                    removed = True
                continue
            buf = ""
        if line.strip(" •-") == flat:
            removed = True
            continue
        out.append(line)
    return "\n".join(out).strip() if removed else t


def load_document(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".pptx":
        return read_pptx(path)
    if ext in (".txt", ".md"):
        return read_text(path)
    return []


def _calm(t: str) -> str:
    """Turn SHOUTING titles into sentence case and tidy stray spacing."""
    letters = [c for c in t if c.isalpha()]
    if letters and sum(c.isupper() for c in letters) / len(letters) > 0.6:
        t = t.lower()
        t = re.sub(r"\bhci\b|\bgoms\b|\bwcag\b|\bsdlc\b|\bui\b|\bux\b", lambda m: m.group(0).upper(), t)
        t = t[:1].upper() + t[1:]
    t = re.sub(r"\s+,", ",", t)
    t = re.sub(r",(?=\S)", ", ", t)
    return re.sub(r"\s+", " ", t).strip()


def doc_title(path: Path, units) -> str:
    """Prefer the module line on the cover page, else the file name."""
    if units:
        head = units[0]["text"]
        m = re.search(r"(Module\s*)(\d+)(\s*[-–—:]*\s*)([^\n]+)", head, re.I)
        if m:
            return f"Module {m.group(2)} – {_calm(m.group(4).strip())}"
    name = re.sub(r"\s*\(\d+\)$", "", path.stem)
    name = re.sub(r"^AIM4141\s*[-–_]*\s*", "", name)
    name = name.replace("_", " ").replace("-", " ").strip()
    return name[:80] or path.stem


def to_sections(units, min_chars=None):
    """Merge units into sections of useful size. Returns list of dict(title, text, refs)."""
    min_chars = min_chars or MIN_SECTION_CHARS
    cleaned = []
    for u in units:
        title = guess_title(u)
        body = strip_title(u, title)
        if BOILERPLATE_TITLE.match(title) or (BOILERPLATE_TEXT.search(u["text"]) and len(u["text"]) < 600):
            continue
        if len(body) < 40:
            continue
        cleaned.append({"title": sentence_case(tidy_title(title)), "text": body, "refs": [u["ref"]]})

    sections = []
    for u in cleaned:
        if sections and len(sections[-1]["text"]) < min_chars:
            prev = sections[-1]
            prev["text"] += f"\n\n{u['title']}\n{u['text']}"
            prev["refs"] += u["refs"]
            if weak_title(prev["title"]) and not weak_title(u["title"]):
                prev["title"] = u["title"]
            elif (not weak_title(u["title"]) and u["title"].lower() not in prev["title"].lower()
                  and len(prev["title"]) + len(u["title"]) < 62):
                prev["title"] = f"{prev['title']} & {u['title']}"
        else:
            sections.append(dict(u))
    # A trailing runt joins the previous section
    if len(sections) > 1 and len(sections[-1]["text"]) < min_chars:
        last = sections.pop()
        sections[-1]["text"] += f"\n\n{last['title']}\n{last['text']}"
        sections[-1]["refs"] += last["refs"]
    return sections


def discover(material_dir: Path):
    files = []
    for p in sorted(material_dir.rglob("*")):
        if p.is_file() and p.suffix.lower() in (".pdf", ".pptx", ".txt", ".md") and not p.name.startswith("."):
            if any(s in p.name.lower() for s in SKIP_PATTERNS):
                continue
            files.append(p)
    return files


def ref_label(refs):
    if not refs:
        return ""
    return refs[0] if len(refs) == 1 else f"{refs[0]} – {refs[-1]}"
